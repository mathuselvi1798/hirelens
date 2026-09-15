"""Raw text extraction from uploaded files.

One function per format, each returning `(text, page_count)`. Any parsing
failure is converted into a domain error here so callers never have to know
which library was used.
"""
import io

from app.core.exceptions import CorruptDocumentError, UnsupportedFileTypeError


def extract_pdf(data: bytes) -> tuple[str, int]:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover
        raise CorruptDocumentError("PDF support is not installed on the server.") from exc

    try:
        parts: list[str] = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                parts.append(page.extract_text() or "")
            page_count = len(pdf.pages)
        return "\n".join(parts), page_count
    except Exception as exc:
        raise CorruptDocumentError(
            "This PDF could not be read. It may be corrupted, password-protected, "
            "or not a valid PDF."
        ) from exc


def extract_docx(data: bytes) -> tuple[str, int | None]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise CorruptDocumentError("DOCX support is not installed on the server.") from exc

    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        # Many resume templates use tables for layout, so pull those in too.
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text and cell.text not in parts:
                        parts.append(cell.text)
        return "\n".join(parts), None
    except Exception as exc:
        raise CorruptDocumentError(
            "This Word document could not be read. It may be corrupted or not a "
            "valid .docx file."
        ) from exc


def extract_txt(data: bytes) -> tuple[str, int | None]:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding), None
        except UnicodeDecodeError:
            continue
    raise CorruptDocumentError("This text file uses an unsupported encoding.")


EXTRACTORS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "txt": extract_txt,
}


def detect_file_type(filename: str) -> str:
    lowered = filename.lower()
    for ext in EXTRACTORS:
        if lowered.endswith(f".{ext}"):
            return ext
    raise UnsupportedFileTypeError(
        f"'{filename}' is not a supported file type. Upload a PDF, DOCX, or TXT file."
    )


def extract(filename: str, data: bytes) -> tuple[str, str, int | None]:
    """Returns `(file_type, text, page_count)`."""
    file_type = detect_file_type(filename)
    text, page_count = EXTRACTORS[file_type](data)
    return file_type, text, page_count
